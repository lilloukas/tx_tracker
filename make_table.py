import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

workbook= openpyxl.load_workbook('notebook.xlsx')


# Convert each sheet to a table
for sheet in workbook.worksheets:
    if sheet.max_row != 1:
        table = Table(displayName=sheet.title.replace(" ", "_"), ref=sheet.dimensions)
        style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        table.tableStyleInfo = style
        # Create the table
        sheet.add_table(table)

# Save the workbook
workbook.save("notebook.xlsx")